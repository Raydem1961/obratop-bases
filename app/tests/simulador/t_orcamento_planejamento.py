import sys,re,json,subprocess,os; sys.path.insert(0,'/tmp/fb')
from harness import *
APP=sys.argv[1];FX='/tmp/fb/fx/';O='organizations/org1/'
res=[]
def chk(n,c,e=''):
    res.append(bool(c));print('OK   ' if c else 'FALHA',n,e,flush=True)
IGN=lambda errs:[e for e in errs if "reading 'update'" not in e and 'xlsx.full.min.js' not in e and '404' not in e and 'Failed to load resource' not in e and 'abrir o PDF' not in e]
def S():
    s=make_seed();s['organizations/org1']={**s['organizations/org1'],'cnpj':'12.345.678/0001-90','companyAddress':'Rua Nova, 10 - Salvador/BA','companyPhone':'(75) 3281-0000','companyEmail':'a@b.com'};return s
def docs(pg,col): return pg.evaluate(f"[...window.__fb.db.entries()].filter(([k])=>k.startsWith('{O}{col}/')&&k.split('/').length===4).map(([k,v])=>({{id:k.split('/')[3],...v}}))")
def go(pg,tab):
    pg.evaluate("document.querySelector('.navbtn[data-route=\"budgets\"]').click()");pg.wait_for_selector('.subtab');pg.click(f'.subtab[data-tab="{tab}"]');pg.wait_for_timeout(500)
def txt(pg,sel): return pg.evaluate("(s)=>{const e=document.querySelector(s);return e?e.textContent:''}",sel)
def brl(s): return float(s.replace('\xa0',' ').replace('R$','').replace('.','').replace(',','.').strip())
a=App(APP,auth=OWNER,seed=S(),port=9800,xlsx='/tmp/vendor/xlsx.full.min.js',viewport={'width':1500,'height':1100})
with a:
    a.open();pg=a.page;pg.set_default_timeout(30000);pg.wait_for_timeout(600);go(pg,'builder');pg.wait_for_selector('#bbName')
    # ---------- etapas antes da busca
    ordem=pg.evaluate("(()=>{const a=document.querySelector('#etList'),b=document.querySelector('#bbQuery');return !!(a.compareDocumentPosition(b)&Node.DOCUMENT_POSITION_FOLLOWING)})()")
    chk('E1 o bloco "Etapas do orçamento" vem ANTES da busca de itens nas bases',ordem and 'Etapas do orçamento' in pg.inner_text('#content'))
    pg.fill('#etNew','Serviços Preliminares');pg.click('#etAdd');pg.fill('#etNew','Fundação');pg.click('#etAdd');pg.fill('#etNew','Paredes');pg.click('#etAdd')
    nomes=pg.eval_on_selector_all('#etList .etRow','e=>e.map(r=>[r.querySelector(".etNo").textContent,r.querySelector(".etName").value])')
    chk('E2 etapas criadas e numeradas 1.0, 2.0, 3.0 (renomeáveis)',nomes==[['1.0','Serviços Preliminares'],['2.0','Fundação'],['3.0','Paredes']],str(nomes))
    pg.click('#etModel');pg.wait_for_timeout(300);n=pg.eval_on_selector_all('#etList .etRow','e=>e.length')
    chk('E3 "etapas padrão de edificação" acrescenta as que faltam sem duplicar (3 existentes + 13 novas)',n==16,str(n))
    for i in range(3,16):   # remove as padrão para deixar só as 3 do teste
        pg.click('#etList .etRow >> nth=3 >> .etDel')
    chk('E4 etapas sem serviços podem ser removidas',pg.eval_on_selector_all('#etList .etRow','e=>e.length')==3)
    pg.set_input_files('#spFile',FX+'SINAPI-2026-08-com-analitico.zip')
    for _ in range(80):
        pg.wait_for_timeout(250)
        if 'composições analíticas' in pg.inner_text('#spStatus') and 'lendo' not in pg.inner_text('#spStatus'): break
    st=pg.inner_text('#spStatus');chk('E5 importar o ZIP traz as composições analíticas (aba Analítico)','3 composições analíticas' in st,st)
    # ---------- adicionar serviços nas etapas
    def add(code,qty,etapa_idx,sub=''):
        pg.select_option('#etActive',index=etapa_idx);pg.fill('#etSub',sub);pg.dispatch_event('#etSub','input');pg.fill('#bbQuery',code);pg.wait_for_timeout(450)
        pg.fill('#bbRes tbody tr >> nth=0 >> .bbQtyAdd',str(qty));pg.click('#bbRes tbody tr >> nth=0 >> .bbAdd');pg.wait_for_timeout(250)
    pg.fill('#bbName','Edificação teste');pg.select_option('#bbWork','w1');pg.fill('#bbBdi','25')
    add('87521',100,2);add('87879',200,2);add('88489',150,1,'Sapatas e cintamento');add('87248',80,1,'Sapatas e cintamento');add('87529',60,1)
    pg.click('#bbManualBox summary');pg.select_option('#etActive',index=0);pg.fill('#etSub','');pg.fill('#mCode','COMP-01');pg.fill('#mDesc','Limpeza inicial do terreno');pg.fill('#mUnit','m2');pg.fill('#mQty','500');pg.fill('#mPrice','2,5');pg.click('#mAdd');pg.wait_for_timeout(250)
    linhas=pg.eval_on_selector_all('#bbTab tbody tr','e=>e.map(r=>[r.className.includes("bbHead")?"H":"I",r.cells[0].innerText.trim(),r.cells[3].innerText.replace(/\\s+/g," ").trim().slice(0,28)])')
    nos=[f'{k}{n}' for k,n,_ in linhas]
    chk('E6 planilha agrupada: etapas 1.0/2.0/3.0, serviço 1.1, serviços e sub-etapa 2.1 (2.1.1, 2.1.2), serviço 2.2 e 3.1/3.2',nos==['H1.0','I1.1','H2.0','H2.1','I2.1.1','I2.1.2','I2.2','H3.0','I3.1','I3.2'],str(nos))
    chk('E7 cada serviço mostra a etapa e a sub-etapa em que está',pg.eval_on_selector_all('#bbTab tbody tr[data-id]','e=>e.map(r=>r.querySelector(".bbSub").value)')[1:3]==['Sapatas e cintamento','Sapatas e cintamento'])
    tag=pg.eval_on_selector_all('#bbTab tbody tr[data-id]','e=>e.map(r=>r.innerText.includes("CPU"))')
    chk('E8 só os serviços com composição analítica no arquivo (87521 e 87879) ganham a marca "CPU"; os demais e o item próprio não',tag.count(True)==2 and tag[0]==False,str(tag))
    # ---------- composição analítica (87521 = linha 2.2? procura pelo código)
    pg.evaluate("[...document.querySelectorAll('#bbTab tbody tr[data-id]')].find(r=>r.cells[2].innerText.trim()==='87521').querySelector('.bbCpu').click()");pg.wait_for_selector('#cpuClose')
    dlg=pg.inner_text('.dialog');lin=pg.eval_on_selector_all('.dialog tbody tr','e=>e.map(r=>[...r.cells].map(c=>c.innerText.trim()))')
    chk('E9 "Composição" mostra a CPU: código, tipo, descrição, unid., coeficiente, preço e custo de cada insumo/composição',len(lin)==5 and lin[0][0]=='00004567' and lin[0][1]=='Insumo' and lin[1][1]=='Composição' and lin[1][0]=='87292' and 'Custo unitário calculado' in dlg,str(lin[:2]))
    chk('E10 mostra o resumo dos insumos explodidos (materiais, mão de obra e equipamento)','material(is)' in dlg and 'equipamento(s)' in dlg and '2 de mão de obra' in dlg,dlg[-260:].replace('\n',' '))
    pg.click('#cpuClose')
    pg.evaluate("[...document.querySelectorAll('#bbTab tbody tr[data-id]')].find(r=>r.cells[2].innerText.trim()==='COMP-01').querySelector('.bbCpu').click()");pg.wait_for_selector('#cpuClose')
    chk('E11 item próprio explica que não tem composição analítica','não tem composição analítica' in pg.inner_text('.dialog'));pg.click('#cpuClose')
    # ---------- totais
    tot=pg.inner_text('#bbTotals').replace('\xa0',' ').replace('\n',' | ');ext=pg.inner_text('#bbExt')
    exp_dir=100*76.8+200*4.6+150*15.7+80*60.3+60*36.4+500*2.5
    chk('E12 totais: custo direto, BDI e valor da obra + valor por extenso',f'{exp_dir:,.2f}'.replace(',','X').replace('.',',').replace('X','.') in tot and 'REAIS' in ext and 'Valor por extenso' in ext,tot+' || '+ext[:90])
    # ---------- exportações
    os.makedirs('/tmp/fb/dl',exist_ok=True)
    with pg.expect_download(timeout=25000) as dl: pg.click('#bbSheetPdf')
    dl.value.save_as('/tmp/fb/dl/planilha.pdf');t=subprocess.run(['pdftotext','-layout','/tmp/fb/dl/planilha.pdf','-'],capture_output=True,text=True).stdout
    chk('E13 PDF da planilha no modelo: cabeçalho da empresa, ITENS/CÓD./DESCRIÇÃO/UNID/QUANT/PR. UNIT/PR. TOTAL, etapas, sub-etapa, TOTAL e valor por extenso',all(x in t for x in ('12.345.678/0001-90','ITENS','DESCRIÇÃO DOS SERVIÇOS','PR. UNIT','PR. TOTAL','Serviços Preliminares','Fundação','Paredes','Sapatas e cintamento','2.1.2','TOTAL','Valor da obra:','REAIS')),re.sub(r'\s+',' ',t)[:240])
    subprocess.run(['pdftoppm','-r','70','-f','1','-l','1','-png','/tmp/fb/dl/planilha.pdf','/tmp/fb/shots/v332_planilha'])
    with pg.expect_download(timeout=25000) as dl2: pg.click('#bbSheetXlsx')
    dl2.value.save_as('/tmp/fb/dl/planilha.xlsx')
    import openpyxl
    ws=openpyxl.load_workbook('/tmp/fb/dl/planilha.xlsx').active;vals=[[c.value for c in r] for r in ws.iter_rows()]
    ult=[r for r in vals if r and r[2]=='TOTAL'][0];it=[r for r in vals if r and r[1]=='87521'][0]
    chk('E14 Excel da planilha no modelo (colunas A–G): preço total é fórmula =QUANT×PR.UNIT e TOTAL é SOMA',vals[0][:7]==['ITENS','CÓD. SINAPI/ORSE','DESCRIÇÃO DOS SERVIÇOS','UNID','QUANT','PR. UNIT ','PR. TOTAL'] and str(it[6]).startswith('=E') and str(ult[6]).startswith('=SUM(G') and any(r and r[0]=='Valor da obra:' for r in vals),f'{it[6]} / {ult[6]}')
    with pg.expect_download(timeout=25000) as dl3: pg.click('#bbCpuPdf')
    dl3.value.save_as('/tmp/fb/dl/cpu.pdf');t3=subprocess.run(['pdftotext','-layout','/tmp/fb/dl/cpu.pdf','-'],capture_output=True,text=True).stdout
    chk('E15 PDF das composições analíticas: cada serviço com insumos, coeficientes, custo unitário e preço com BDI; item próprio avisa que não tem CPU',all(x in t3 for x in ('87521','00004567','PEDREIRO COM ENCARGOS','Custo unitário','Preço unitário com BDI 25%','Sem composição analítica')),re.sub(r'\s+',' ',t3)[:200])
    with pg.expect_download(timeout=25000) as dl4: pg.click('#bbCpuXlsx')
    dl4.value.save_as('/tmp/fb/dl/cpu.xlsx');cw=openpyxl.load_workbook('/tmp/fb/dl/cpu.xlsx').active;cv=[[c.value for c in r] for r in cw.iter_rows()]
    chk('E16 Excel das composições analíticas com os insumos de cada serviço',cv[0][:3]==['ITEM','CÓDIGO','TIPO'] and sum(1 for r in cv if r and r[2]=='Insumo')>=2 and any(r and r[2]=='Composição' for r in cv))
    # ---------- remoção de etapa com serviços e ordem
    n0=pg.eval_on_selector_all('#etList .etRow','e=>e.length');m=len(a.dialogs);pg.click('#etList .etRow >> nth=1 >> .etDel');pg.wait_for_timeout(200)
    chk('E17 etapa com serviços não pode ser excluída (aviso)',pg.eval_on_selector_all('#etList .etRow','e=>e.length')==n0 and any('tem serviços' in x for _,x in a.dialogs[m:]))
    # ---------- rascunho com etapas
    pg.reload();pg.wait_for_timeout(1500);go(pg,'builder');pg.wait_for_timeout(600)
    chk('E18 o rascunho guarda etapas, serviços e a composição analítica depois de recarregar',pg.eval_on_selector_all('#etList .etRow','e=>e.length')==3 and pg.eval_on_selector_all('#bbTab tbody tr[data-id]','e=>e.length')==6 and pg.eval_on_selector_all('#bbTab tbody tr[data-id]','e=>e.filter(r=>r.innerText.includes("CPU")).length')==2)
    # ---------- salvar e gerar planejamento
    pg.select_option('#bbWork','w1');m=len(a.dialogs);pg.click('#bbPlan');pg.wait_for_selector('#pgOk')
    pg.fill('#pgStart','2026-11-02');pg.fill('#pgEnd','2027-03-31');pg.click('#pgOk');pg.wait_for_timeout(2500)
    bd=[x for x in docs(pg,'budgets') if x.get('budgetName')=='Edificação teste'];ac=[x for x in docs(pg,'activities') if x.get('budgetId') in {b['id'] for b in bd}]
    chk('E19 o orçamento é gravado com etapa, número do item, composição (insumos por tipo) e CPU',len(bd)==6 and all(b.get('etapaNo') and b.get('itemNo') and b['workId']=='w1' for b in bd) and sum(1 for b in bd if b.get('composition'))==2 and sum(1 for b in bd if b.get('cpu'))==2,str(sorted((b['itemNo'],b['etapa']) for b in bd))[:200])
    b87=[b for b in bd if b['sourceCode']=='87521'][0]
    chk('E20 o item 87521 guarda os insumos explodidos (bloco, cimento, areia, servente, pedreiro, betoneira) com coeficiente e tipo',{c['kind'] for c in b87['composition']}=={'Material','Mão de obra','Equipamento'} and any(c['code']=='00000001' and abs(c['coef']-0.16)<1e-9 for c in b87['composition']) and len(b87['cpu'])==5,str([(c['code'],c['kind'],c['coef']) for c in b87['composition']]))
    chk('E21 gera uma atividade por serviço, ligada ao item do orçamento, com EAP igual à numeração da planilha, datas no prazo e dias úteis',len(ac)==6 and all(x['start']>='2026-11-02' and x['end']<='2027-03-31' and x['durationDays']>=1 for x in ac) and {x['wbs'] for x in ac}=={b['itemNo'] for b in bd},str(sorted((x['wbs'],x['start'],x['end']) for x in ac)))
    chk('E22 as atividades levam os grupos da EAP por etapa (nome e código da etapa) e predecessoras',all(x.get('eapPhaseName') and x.get('eapPhaseCode','').endswith('.0') for x in ac) and any(x['predecessors'] for x in ac))
    w=[x for x in docs(pg,'works') if x['id']=='w1'][0]
    chk('E23 cria a linha de base R1 do cronograma',w.get('baseline',{}).get('revision')==1 and len(w['baseline']['items'])==6)
    chk('E24 depois de gerar, abre a aba Planejamento e quantitativos da obra',pg.eval_on_selector('.subtab.active','e=>e.dataset.tab')=='plan' and pg.input_value('#plWork')=='w1')
    chk('E25 sem erros de JavaScript (etapas, CPU, exportações e geração)',not IGN(a.errors),str(IGN(a.errors)[:2]))
    # ---------- PLANEJAMENTO
    eap=pg.inner_text('#plBody');chk('P1 EAP: etapas e serviços com início, término e valor; cartões de resumo',all(x in eap for x in ('Serviços Preliminares','Fundação','Paredes','2.1.1','Valor da obra (c/ BDI)','Início → término')) and pg.eval_on_selector_all('#plBody tr.plEt','e=>e.length')==3,eap[:200].replace('\n',' | ')+' | plEt='+str(pg.eval_on_selector_all('#plBody tr.plEt','e=>e.length')))
    pg.locator('#plRoot').screenshot(path='/tmp/fb/shots/v332_eap.png');chk('P2 o serviço sem composição (item próprio) aparece sinalizado','sem CPU' in eap)
    pg.click('.plSegBtn[data-v="curva"]');pg.wait_for_timeout(300)
    meses=pg.eval_on_selector_all('#plBody tbody tr','e=>e.map(r=>r.cells[0].innerText)')
    chk('P3 Curva S: gráfico previsto x realizado e tabela mês a mês até 100% previsto',pg.query_selector('#plBody svg polyline') is not None and len(meses)>=4 and pg.evaluate("(()=>{const h=[...document.querySelectorAll('#plBody thead th')].findIndex(x=>x.textContent.trim()==='% previsto');return document.querySelector('#plBody tbody tr:last-child').cells[h].innerText})()").startswith('100'),str(meses)+' '+' svg='+str(pg.query_selector('#plBody svg polyline') is not None))
    pg.locator('#plRoot').screenshot(path='/tmp/fb/shots/v332_curva.png');pg.click('.plSegBtn[data-v="abc"]');pg.wait_for_timeout(300);pg.locator('#plRoot').screenshot(path='/tmp/fb/shots/v332_abc.png');abc=pg.inner_text('#plBody')
    chk('P4 Curvas ABC de materiais, mão de obra e equipamentos, com classe A/B/C',all(x in abc for x in ('Curva ABC — Materiais','Curva ABC — Mão de obra','Curva ABC — Equipamentos','Classe A')) and 'BETONEIRA' in abc and 'SERVENTE' in abc and 'BLOCO' in abc)
    pg.click('.plSegBtn[data-v="Material"]');pg.wait_for_timeout(300);mat=pg.inner_text('#plBody')
    chk('P5 Quantitativo de materiais consolidado: bloco = 13 × 100 m² = 1.300 un; cimento = 16 (argamassa) + 10 (chapisco) = 26 sc',re.search(r'BLOCO[^\n]*\t[^\n]*1\.300,00',mat) is not None and re.search(r'CIMENTO[^\n]*\t[^\n]*26,00',mat) is not None,mat[:300].replace('\n',' | '))
    pg.check('#plMonthly');pg.wait_for_timeout(300);colunas=pg.eval_on_selector_all('#plBody thead th','e=>e.length')
    chk('P6 "Mês a mês" abre as colunas previsto e realizado de cada mês',colunas>=9+2*4 and 'previsto' in pg.inner_text('#plBody thead') and 'realizado' in pg.inner_text('#plBody thead'),str(colunas))
    pg.click('.plSegBtn[data-v="Mão de obra"]');pg.wait_for_timeout(300);mo=pg.inner_text('#plBody')
    chk('P7 Mão de obra em horas (servente 60 h, pedreiro 80 h) com o efetivo médio por mês','SERVENTE' in mo and re.search(r'PEDREIRO[^\n]*80,00',mo) is not None and 'Efetivo médio' in mo,mo[:200].replace('\n',' | '))
    pg.locator('#plRoot').screenshot(path='/tmp/fb/shots/v332_mo.png');pg.click('.plSegBtn[data-v="Equipamento"]');pg.wait_for_timeout(300);chk('P8 Equipamentos: betoneira 1 CHP','BETONEIRA' in pg.inner_text('#plBody'))
    pg.check('input[name=plMode][value=servico]');pg.wait_for_timeout(300)
    chk('P9 Quantitativo por serviço: o insumo mostra coeficiente × quantidade do serviço','ALVENARIA' in pg.inner_text('#plBody') and 'Coeficiente' in pg.inner_text('#plBody'),pg.inner_text('#plBody')[:200].replace('\n',' | '))
    with pg.expect_download(timeout=25000) as dl5: pg.click('#plXlsx')
    dl5.value.save_as('/tmp/fb/dl/plan_eq.xlsx');wb=openpyxl.load_workbook('/tmp/fb/dl/plan_eq.xlsx')
    chk('P10 Excel do quantitativo traz o consolidado (com meses) e o detalhe por serviço',wb.sheetnames==['Consolidado','Por serviço'] and wb['Consolidado'].cell(1,1).value=='Código')
    # ---------- realizado
    pg.click('.plSegBtn[data-v="real"]');pg.wait_for_timeout(300);mes=pg.input_value('#plMonth');chk('P11 "Lançar realizado" lista os serviços previstos no mês',pg.eval_on_selector_all('#plBody .plRealIn','e=>e.length')>=1,mes)
    pg.locator('#plRoot').screenshot(path='/tmp/fb/shots/v332_real.png');pg.click('#plFill');pg.wait_for_timeout(200);pg.click('#plSaveReal');pg.wait_for_timeout(1500)
    bd2=[x for x in docs(pg,'budgets') if x.get('budgetName')=='Edificação teste'];com=[b for b in bd2 if (b.get('real') or {}).get(mes)]
    chk('P12 salvar grava a quantidade executada do mês em cada serviço (campo real)',len(com)>=1 and all(isinstance(b['real'][mes],(int,float)) for b in com),str([(b['itemNo'],b['real']) for b in com][:3]))
    pg.click('.plSegBtn[data-v="curva"]');pg.wait_for_timeout(400);r1=pg.eval_on_selector('#plBody tbody tr:first-child','r=>[r.cells[2].innerText,r.cells[6].innerText]')
    chk('P13 a Curva S passa a mostrar o realizado do mês lançado (R$ e %)',brl(r1[0])>0 and r1[1]!='0%',str(r1))
    pg.click('.plSegBtn[data-v="Material"]');pg.wait_for_timeout(300)
    real_q=pg.eval_on_selector_all('#plBody tbody tr','e=>e.map(r=>[r.cells[1].innerText,r.cells[3].innerText,r.cells[4].innerText])')
    chk('P14 o quantitativo realizado de materiais é calculado pelos coeficientes (realizado > 0 depois do lançamento)',any(r[0].startswith('BLOCO') and r[2]!='0,00' for r in real_q),str(real_q[:3]))
    # ---------- Cronograma (Gantt)
    pg.evaluate("document.querySelector('.navbtn[data-route=\"activities\"]').click()");pg.wait_for_timeout(900);g=pg.inner_text('#content')
    chk('P15 o Cronograma (Gantt) mostra a EAP do orçamento agrupada por etapa, com os códigos do orçamento','Serviços Preliminares' in g and 'Fundação' in g and '2.1.1' in g and '1.0' in g,g[:240].replace('\n',' | '))
    pg.screenshot(path='/tmp/fb/shots/v332_gantt.png',clip={'x':220,'y':80,'width':1280,'height':620})
    chk('P16 sem erros de JavaScript (planejamento)',not IGN(a.errors),str(IGN(a.errors)[:2]))
# ---------- perfil Consulta e celular
def mem(uid,role,wid):
    s=S();s[O+'members/'+uid]={'email':uid+'@t.com','name':uid,'role':role,'status':'active','workId':wid};s['users/'+uid]={'currentOrg':'org1','email':uid+'@t.com'};return s
with App(APP,auth={'uid':'cons','email':'cons@t.com','emailVerified':True},seed=mem('cons','viewer','w1'),port=9801,viewport={'width':1500,'height':1100}) as d_:
    d_.open();pg=d_.page;go(pg,'builder');pg.wait_for_timeout(400)
    chk('V1 perfil Consulta monta e exporta, sem os botões de gravar e de gerar planejamento',pg.is_hidden('#bbSave') and pg.is_hidden('#bbPlan') and pg.is_visible('#bbSheetPdf'))
    go(pg,'plan');pg.wait_for_timeout(500);chk('V2 e vê o Planejamento (sem permissão para lançar realizado)','Esta obra ainda não tem orçamento' in pg.inner_text('#content') or pg.query_selector('#plSaveReal') is None)
m=App(APP,auth=OWNER,seed=S(),port=9802,xlsx='/tmp/vendor/xlsx.full.min.js',viewport={'width':390,'height':844},mobile=True)
with m:
    m.open();pg=m.page;pg.wait_for_timeout(600);pg.evaluate("document.querySelector('.navbtn[data-route=\"budgets\"]').click()");pg.wait_for_selector('.subtab');pg.tap('.subtab[data-tab="builder"]');pg.wait_for_selector('#bbName')
    pg.fill('#etNew','Fundação');pg.tap('#etAdd');pg.tap('#bbManualBox summary');pg.fill('#mDesc','Item próprio');pg.fill('#mPrice','10');pg.tap('#mAdd');pg.wait_for_timeout(300)
    ov=pg.evaluate("Math.max(document.documentElement.scrollWidth,document.body.scrollWidth)-Math.min(document.documentElement.clientWidth,screen.width)")
    chk('M1 no celular a tela com etapas e planilha agrupada não estoura a largura',ov<=1,str(ov));pg.screenshot(path='/tmp/fb/shots/v332_mobile.png')
    pg.evaluate("document.querySelector('.subtab[data-tab=\"plan\"]').click()");pg.wait_for_timeout(500);chk('M2 e a aba Planejamento abre no celular sem estourar a largura',pg.evaluate("Math.max(document.documentElement.scrollWidth,document.body.scrollWidth)-Math.min(document.documentElement.clientWidth,screen.width)")<=1)
print('\nRESUMO:',sum(res),'de',len(res));sys.exit(0 if all(res) else 1)
