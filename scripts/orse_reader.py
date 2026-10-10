"""Leitor da base ORSE (CEHOP/SE) a partir da planilha exportada do programa ORSE.
Abas usadas: Leia-me (período), Composições (preço unitário) e Insumos (preço adotado)."""
from __future__ import annotations
import re, pathlib
import openpyxl

def _num(v):
    if isinstance(v,(int,float)): return round(float(v),4)
    return None

def read_orse_xlsx(path):
    """Retorna (referencia 'AAAA-MM', linhas). Cada linha: code, description, unit, price, sheet, group."""
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    names=set(wb.sheetnames)
    for need in ("Composições","Insumos"):
        if need not in names: raise ValueError(f"planilha ORSE sem a aba '{need}'")
    ref=None
    if "Leia-me" in names:
        for r in wb["Leia-me"].iter_rows(values_only=True):
            if r and r[0] and "per" in str(r[0]).lower() and r[1]:
                m=re.search(r"(\d{1,2})\s*/\s*(20\d{2})",str(r[1]))
                if m: ref=f"{m.group(2)}-{int(m.group(1)):02d}"
    if not ref:
        m=re.search(r"(20\d{2})[-_]?(0[1-9]|1[0-2])",pathlib.Path(path).name)
        if m: ref=f"{m.group(1)}-{m.group(2)}"
    if not ref: raise ValueError("não foi possível identificar o período da planilha ORSE")
    rows=[]
    def scan(sheet,pcol,label):
        it=wb[sheet].iter_rows(values_only=True); hdr=next(it)
        if str(hdr[1]).strip()!="Código" or "Descri" not in str(hdr[2]): raise ValueError(f"cabeçalho inesperado em {sheet}")
        for r in it:
            if not r or r[1] in (None,""): continue
            price=_num(r[pcol]); desc=str(r[2] or "").strip(); unit=str(r[3] or "").strip()
            if not price or price<=0 or not desc or not unit: continue
            code=r[1]; code=str(int(code)) if isinstance(code,(int,float)) else str(code).strip()
            rows.append({"code":code,"description":desc,"unit":unit,"price":price,"sheet":label,"group":str(r[4] or "").strip()})
    scan("Composições",6,"Composições"); scan("Insumos",7,"Insumos")
    return ref,rows
