"""Leitor do relatório mensal SINAPI (CAIXA, "Relatórios a partir de 2025", formato XLSX).

Arquivo lido: SINAPI_Referência_AAAA_MM.xlsx
  ISD / ICD : preços medianos de INSUMOS (sem / com desoneração), UFs em colunas
  CSD / CCD : custos de COMPOSIÇÕES (sem / com desoneração), UFs em colunas (Custo, %AS)
As abas ISE / CSE (sem encargos) não são usadas.
Preço vazio = sem preço/custo para a UF: o item é ignorado (nunca vira zero).
"""
from __future__ import annotations
import re, unicodedata
import openpyxl

UFS = ["AC","AL","AM","AP","BA","CE","DF","ES","GO","MA","MG","MS","MT","PA","PB","PE","PI","PR","RJ","RN","RO","RR","RS","SC","SE","SP","TO"]
REGIMES = {"ISD": ("Não desonerado", "insumo"), "CSD": ("Não desonerado", "composicao"),
           "ICD": ("Desonerado", "insumo"), "CCD": ("Desonerado", "composicao")}

def _n(v):
    s = unicodedata.normalize("NFD", str(v or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()

def _price(v):
    if v is None or v == "": return 0.0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip().replace("R$", "").replace(" ", "")
    if not s: return 0.0
    if "," in s and "." in s: s = s.replace(".", "").replace(",", ".")
    elif "," in s: s = s.replace(",", ".")
    try: return float(s)
    except ValueError: return 0.0

def _text(v):
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()

def _code(v):
    s = _text(v)
    if re.fullmatch(r"\d+\.0", s): s = s[:-2]
    return s

def _find_header(rows):
    """Linha do cabeçalho: contém 'codigo' e 'descricao' e 'unidade'."""
    for i, r in enumerate(rows[:30]):
        t = [_n(x) for x in r]
        if any(x.startswith("codigo") for x in t) and any(x.startswith("descricao") for x in t) and any(x.startswith("unidade") for x in t):
            return i
    return None

def _composition_codes(wb):
    """(grupo, descrição, unidade) -> código, lido da aba Analítico (a coluna de código das abas de custo é uma fórmula sem valor guardado)."""
    m = {}
    if "Analítico" not in wb.sheetnames: return m
    for r in wb["Analítico"].iter_rows(min_row=11, values_only=True):
        if r and len(r) > 5 and r[1] and r[2] is None and r[3] is None and r[4]:
            m[(_text(r[0]), _text(r[4]), _text(r[5]))] = _code(r[1])
    return m

def _read_sheet(ws, kind, codes=None):
    rows = list(ws.iter_rows(values_only=True))
    h = _find_header(rows)
    if h is None: raise ValueError(f"cabeçalho não encontrado na aba {ws.title}")
    head = [_n(x) for x in rows[h]]
    ci = next(i for i, x in enumerate(head) if x.startswith("codigo"))
    di = next(i for i, x in enumerate(head) if x.startswith("descricao"))
    ui = next(i for i, x in enumerate(head) if x.startswith("unidade"))
    # colunas de preço por UF
    cols = {}
    if kind == "insumo":
        for j, x in enumerate(rows[h]):
            u = _text(x).upper()
            if u in UFS and u not in cols: cols[u] = j
    else:
        # composições: UFs na linha acima do cabeçalho (cada UF ocupa 2 colunas: Custo, %AS)
        up = rows[h - 1]
        for j, x in enumerate(up):
            u = _text(x).upper()
            if u in UFS and u not in cols and _n(rows[h][j]).startswith("custo"): cols[u] = j
    if len(cols) < 20: raise ValueError(f"aba {ws.title}: só {len(cols)} UFs encontradas")
    out = {u: [] for u in cols}
    gi = 0  # classificação / grupo
    for r in rows[h + 1:]:
        if r is None or len(r) <= max(cols.values()): continue
        code, desc, unit = _code(r[ci]), _text(r[di]), _text(r[ui])
        grp = _text(r[gi])
        if kind == "composicao" and codes is not None: code = codes.get((grp, desc, unit), "")
        if not desc or not unit or code in ("", "0"): continue
        for u, j in cols.items():
            p = _price(r[j])
            if p > 0 and p < 1e9:
                out[u].append({"code": code, "description": desc, "unit": unit, "price": round(p, 2), "sheet": ws.title, "group": grp})
    return out

def read_sinapi_reference(path):
    """Devolve {(uf, regime): [linhas]} juntando insumos e composições; e o mês de referência."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ref = ""
    result = {}
    try:
        codes = _composition_codes(wb)
        for sheet, (regime, kind) in REGIMES.items():
            if sheet not in wb.sheetnames: continue
            ws = wb[sheet]
            if not ref:
                for r in ws.iter_rows(min_row=1, max_row=6, values_only=True):
                    if r and _n(r[0]).startswith("mes de referencia"):
                        m = re.search(r"(0[1-9]|1[0-2])\s*/\s*(20\d{2})", _text(r[1])); ref = f"{m.group(2)}-{m.group(1)}" if m else ""
            for uf, rows in _read_sheet(ws, kind, codes).items():
                result.setdefault((uf, regime), []).extend(rows)
    finally:
        wb.close()
    return ref, result
